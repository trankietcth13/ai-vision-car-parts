from openpyxl import load_workbook

from image_db_crawler.models import ClassificationResult, VehicleMetadata
from image_db_crawler.storage import DatasetStore


def test_image_and_excel_are_stored_together(tmp_path):
    store = DatasetStore(tmp_path)
    image = store.save(
        b"unique image bytes",
        ".jpg",
        VehicleMetadata("2014", "Falcon", "F7", "7.0L V8"),
        ClassificationResult(True, "engine_bay", "engine_cover", 0.91, ("engine_cover",)),
        "https://media.carsandbids.com/x/photos/1.jpg",
        "https://carsandbids.com/auctions/x/car",
        source_name="wikimedia-commons",
        license_code="cc-by-sa-4.0",
        creator="Example Author",
    )
    duplicate = store.save(
        b"unique image bytes", ".jpg", VehicleMetadata(), ClassificationResult(True), "x", "y"
    )
    store.close()
    assert image is not None and image.exists()
    workbook_path = image.parent / "images.xlsx"
    assert workbook_path.exists()
    row = list(load_workbook(workbook_path).active.iter_rows(values_only=True))[1]
    assert row[0] == 1
    assert row[1] == "2014_Falcon_F7_7.0L_V8_00000001"
    assert row[6] == "engine_cover"
    assert row[10] == "wikimedia-commons"
    assert row[13] == "cc-by-sa-4.0"
    assert row[15] == "Example Author"
    assert duplicate is None
