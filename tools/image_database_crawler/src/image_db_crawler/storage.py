from __future__ import annotations

import hashlib
import re
import sqlite3
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .models import ClassificationResult, VehicleMetadata


HEADERS = [
    "id", "ymme_id", "year", "make", "model", "engine", "component_name", "category",
    "confidence", "image_file", "source_name", "source_image_url", "source_page_url",
    "license_code", "license_url", "creator", "creator_url", "attribution", "component_hint",
    "sha256", "collected_at_utc",
]


def safe_segment(value: str, fallback: str) -> str:
    value = re.sub(r"[<>:\"/\\|?*\x00-\x1f]", "_", value).strip(" ._")
    value = re.sub(r"\s+", "_", value)
    return value[:100] or fallback


class DatasetStore:
    def __init__(self, root: Path) -> None:
        self.root = root.resolve()
        self.root.mkdir(parents=True, exist_ok=True)
        state = self.root / ".crawler"
        state.mkdir(exist_ok=True)
        self.db = sqlite3.connect(state / "catalog.sqlite3")
        self.db.execute("CREATE TABLE IF NOT EXISTS ids (id INTEGER PRIMARY KEY AUTOINCREMENT, sha256 TEXT UNIQUE NOT NULL)")
        self.db.commit()

    def close(self) -> None:
        self.db.close()

    def save(
        self,
        content: bytes,
        extension: str,
        vehicle: VehicleMetadata,
        classification: ClassificationResult,
        source_url: str,
        listing_url: str,
        source_name: str = "",
        license_code: str = "",
        license_url: str = "",
        creator: str = "",
        creator_url: str = "",
        attribution: str = "",
        component_hint: str = "",
    ) -> Path | None:
        digest = hashlib.sha256(content).hexdigest()
        try:
            cursor = self.db.execute("INSERT INTO ids (sha256) VALUES (?)", (digest,))
            self.db.commit()
        except sqlite3.IntegrityError:
            return None
        image_id = int(cursor.lastrowid)
        parts = [
            safe_segment(vehicle.year, "unknown-year"),
            safe_segment(vehicle.make, "unknown-make"),
            safe_segment(vehicle.model, "unknown-model"),
            safe_segment(vehicle.engine, "unknown-engine"),
        ]
        folder = self.root.joinpath(*parts)
        folder.mkdir(parents=True, exist_ok=True)
        filename = f"{image_id:08d}{extension}"
        image_path = folder / filename
        image_path.write_bytes(content)
        ymme_id = "_".join(parts + [f"{image_id:08d}"])
        row = [
            image_id, ymme_id, vehicle.year, vehicle.make, vehicle.model, vehicle.engine,
            classification.component_name, classification.category, classification.confidence,
            filename, source_name, source_url, listing_url, license_code, license_url, creator,
            creator_url, attribution, component_hint, digest, datetime.now(timezone.utc).isoformat(),
        ]
        self._append_excel(folder / "images.xlsx", row)
        return image_path

    @staticmethod
    def _append_excel(path: Path, row: list[object]) -> None:
        if path.exists():
            workbook = load_workbook(path)
            sheet = workbook.active
        else:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "images"
            sheet.append(HEADERS)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = f"A1:U1"
        sheet.append(row)
        widths = {
            "A": 10, "B": 55, "C": 12, "D": 20, "E": 28, "F": 28, "G": 45,
            "H": 20, "I": 12, "J": 18, "K": 22, "L": 60, "M": 60, "N": 22,
            "O": 48, "P": 28, "Q": 48, "R": 60, "S": 45, "T": 68, "U": 28,
        }
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width
        workbook.save(path)
